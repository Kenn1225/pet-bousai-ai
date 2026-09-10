# -*- coding: utf-8 -*-
"""
PDPA 公式スライドテンプレート（共用モジュール）

一般社団法人ペット防災アドバイザー協会の教材デザインルールを実装したもの。
今後の教材（PDPA-101〜112、D201/C202、I401、R501）はこのモジュールを import して
中身だけを書く。デザインを変えるときはこの1ファイルだけを直す。

協会標準（固定・変更しない）
  ブランドカラー : ネイビー #17375E / オレンジ #F5A623 / ホワイト #FFFFFF / ライトグレー #F3F5F7
  フォント       : BIZ UDPゴシック
  文字サイズ     : タイトル 34〜40pt / 本文 26〜30pt / 注釈 18〜20pt
  デザインルール : タイトル=ネイビー帯白文字 / 見出し=オレンジ帯 / 重要=赤い「！」
                   ワーク=緑帯 / ケーススタディ=オレンジ背景 / コラム=薄いグレー
                   1ページ1テーマ / 本文80文字程度 / 写真必須 / 図必須 / 出典必須(下部18pt)

実装上の注意（実測で判明した落とし穴）
  1. PowerPoint の BIZ UDPGothic の実効行高は
     「フォントpt × line_spacing × 約1.19 ÷ 72」インチ。
     本文26pt なら line_spacing は 1.2 が安全上限。これを超えると下にはみ出す。
  2. python-pptx は a:latin しか設定しないため、日本語は a:ea / a:cs も指定が必要。
     子要素の順序は latin → ea → cs（違えると PowerPoint が修復ダイアログを出す）。
"""
import json
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from PIL import Image

# ---------------- 公式ブランドカラー ----------------
NAVY = RGBColor(0x17, 0x37, 0x5E)
ORANGE = RGBColor(0xF5, 0xA6, 0x23)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY = RGBColor(0xF3, 0xF5, 0xF7)
# 派生色（公式4色から演算・用途限定）
NAVY_CARD = RGBColor(0x24, 0x48, 0x70)
ORANGE_TXT = RGBColor(0xB3, 0x76, 0x08)
TEXT = RGBColor(0x1B, 0x22, 0x2C)
MUTED = RGBColor(0x5A, 0x64, 0x72)
BORDER = RGBColor(0xD8, 0xDE, 0xE6)
RED = RGBColor(0xC0, 0x2A, 0x1E)      # 重要「！」専用
GREEN = RGBColor(0x2E, 0x7D, 0x5B)    # ワーク帯専用
CREAM_MUTED = RGBColor(0xC9, 0xD2, 0xDD)

FONT = "BIZ UDPGothic"
# 認定証など「権威性を出す印刷物」の作図にだけ使う明朝体。
# スライド本文には使わない（協会標準は BIZ UDPゴシック）。
MINCHO = "BIZ UDPMincho Medium"

# ---------------- 文字サイズ（協会標準） ----------------
SZ_TITLE = 34
SZ_TITLE_L = 40
SZ_BODY = 26
SZ_BODY_L = 30
SZ_NOTE = 18
SZ_NOTE_L = 20

# ---------------- レイアウト定数（16:9 / 13.333×7.5in） ----------------
SW = Inches(13.333)
SH = Inches(7.5)
MARGIN = Inches(0.75)
CW = SW - MARGIN * 2          # コンテンツ幅 11.833in
BAND_H = Inches(1.30)         # タイトル帯
CY = Inches(1.60)             # コンテンツ上端
CB = Inches(6.45)             # コンテンツ下端（出典帯に触れない限界）
CH = CB - CY
SRC_Y = Inches(6.55)
SRC_H = Inches(0.95)

# 本文サイズごとの安全な line_spacing 上限（上記の落とし穴1より）
SAFE_LS = {30: 1.15, 28: 1.18, 26: 1.20, 20: 1.30, 18: 1.30}


def safe_ls(size):
    return SAFE_LS.get(int(size), 1.25)


def _apply_ea(run, font_name=FONT):
    """日本語（East Asian）フォントも同じ書体に固定する。順序は latin → ea → cs"""
    rPr = run._r.get_or_add_rPr()
    anchor = rPr.find(qn("a:latin"))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = parse_xml(
                '<%s xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>' % tag)
            if anchor is not None:
                anchor.addnext(el)
            else:
                rPr.append(el)
        el.set("typeface", font_name)
        anchor = el


class Deck:
    """1教材ぶんのスライド束。

    code        : 教材コード（PDPA-100S 等）。slide_refs.json に記録される
    photo_dirs  : 写真を探すフォルダ。先に指定したものが優先される
    ref_map     : {ページ番号: [REF番号, ...]}。出典管理台帳との紐付け
    """

    def __init__(self, code, photo_dirs=(), ref_map=None, version="1.0"):
        self.code = code
        self.version = version
        self.photo_dirs = [d for d in photo_dirs if os.path.isdir(d)]
        self.ref_map = ref_map or {}
        self.slide_refs = {}
        self.page = 0

        self.prs = Presentation()
        self.prs.slide_width = SW
        self.prs.slide_height = SH
        self._blank = self.prs.slide_layouts[6]

    # ---------- 素材 ----------
    def photo(self, name):
        for d in self.photo_dirs:
            for ext in (".jpg", ".jpeg", ".png"):
                p = os.path.join(d, name + ext)
                if os.path.exists(p):
                    return p
        raise FileNotFoundError("photo not found: %s (searched %s)"
                                % (name, self.photo_dirs))

    # ---------- 基本図形 ----------
    def add_slide(self, bg=WHITE):
        s = self.prs.slides.add_slide(self._blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = bg
        return s

    def rect(self, slide, x, y, w, h, fill=None, line=None, radius=None, line_w=1.25):
        st = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shp = slide.shapes.add_shape(st, int(x), int(y), int(w), int(h))
        if radius:
            try:
                shp.adjustments[0] = radius
            except Exception:
                pass
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = Pt(line_w)
        shp.shadow.inherit = False
        shp.text_frame.word_wrap = True
        return shp

    def tbox(self, slide, x, y, w, h, text, size=SZ_BODY, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=None, italic=False,
             font=None, space=None):
        if ls is None:
            ls = safe_ls(size)
        font = font or FONT
        tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = ls
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.name = font
            r.font.color.rgb = color
            _apply_ea(r, font)
            if space:   # 字間（認定証の「認 定 証」など）。1/100pt 単位
                r._r.get_or_add_rPr().set("spc", str(int(space * 100)))
        return tb

    def pic_fill(self, slide, name, x, y, w, h):
        """指定枠を埋めるようにセンタークロップして写真を配置"""
        path = self.photo(name)
        x, y, w, h = int(x), int(y), int(w), int(h)
        iw, ih = Image.open(path).size
        pic = slide.shapes.add_picture(path, x, y, width=w, height=h)
        tr, sr = w / h, iw / ih
        if sr > tr:
            c = (1 - tr / sr) / 2
            pic.crop_left = pic.crop_right = c
        else:
            c = (1 - sr / tr) / 2
            pic.crop_top = pic.crop_bottom = c
        return pic

    def framed_pic(self, slide, name, x, y, w, h, frame=WHITE, fw=3.0):
        self.pic_fill(slide, name, x, y, w, h)
        return self.rect(slide, x, y, w, h, fill=None, line=frame, line_w=fw)

    @staticmethod
    def translucent(shape, transparency_pct):
        """図形に透明度を与える（0=不透明, 100=完全透明）"""
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        srgb = shape.fill._xPr.find(".//" + ns + "srgbClr")
        el = srgb.makeelement(ns + "alpha",
                              {"val": str(int((100 - transparency_pct) * 1000))})
        srgb.append(el)

    # ---------- 協会デザインルールの部品 ----------
    def title_band(self, slide, text, band=NAVY, txt=WHITE, size=SZ_TITLE,
                   kicker=None, kicker_color=None):
        """タイトル＝ネイビー帯＋白文字。下端にオレンジのアクセント線"""
        self.rect(slide, 0, 0, SW, BAND_H, fill=band)
        if kicker:
            self.tbox(slide, MARGIN, Inches(0.16), CW, Inches(0.32), kicker,
                      size=SZ_NOTE, color=kicker_color or ORANGE, bold=True)
            self.tbox(slide, MARGIN, Inches(0.52), CW, Inches(0.66), text,
                      size=size, color=txt, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        else:
            self.tbox(slide, MARGIN, 0, CW, BAND_H, text, size=size, color=txt,
                      bold=True, anchor=MSO_ANCHOR.MIDDLE)
        self.rect(slide, 0, BAND_H - Pt(5), SW, Pt(5), fill=ORANGE)

    def orange_head(self, slide, x, y, w, text, h=Inches(0.52), size=SZ_NOTE_L):
        """見出し＝オレンジ帯"""
        self.rect(slide, x, y, w, h, fill=ORANGE)
        self.tbox(slide, x + Inches(0.18), y, w - Inches(0.36), h, text,
                  size=size, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    def alert_icon(self, slide, x, y, d=Inches(0.62)):
        """重要＝赤い「！」アイコン"""
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(x), int(y), int(d), int(d))
        c.fill.solid(); c.fill.fore_color.rgb = RED
        c.line.fill.background(); c.shadow.inherit = False
        tf = c.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = "！"
        r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = WHITE
        r.font.name = FONT; _apply_ea(r)
        return c

    def num_badge(self, slide, x, y, n, d=Inches(0.56), fill=ORANGE, color=NAVY):
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(x), int(y), int(d), int(d))
        c.fill.solid(); c.fill.fore_color.rgb = fill
        c.line.fill.background(); c.shadow.inherit = False
        tf = c.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False          # 2桁の番号が円内で折り返さないように
        tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(n)
        r.font.size = Pt(20 if len(str(n)) < 2 else 15)
        r.font.bold = True; r.font.color.rgb = color
        r.font.name = FONT; _apply_ea(r)
        return c

    def source_band(self, slide, source, dark=False, page=True):
        """出典＝ページ下部・18pt（協会標準）。あわせて REF 番号を記録する"""
        self.page += 1
        n = self.page
        self.slide_refs[n] = {"source": source, "refs": self.ref_map.get(n, [])}
        bg = NAVY_CARD if dark else LGRAY
        fg = CREAM_MUTED if dark else MUTED
        self.rect(slide, 0, SRC_Y, SW, SRC_H, fill=bg)
        self.tbox(slide, MARGIN, SRC_Y, CW - Inches(1.1), SRC_H, source,
                  size=SZ_NOTE, color=fg, ls=1.15, anchor=MSO_ANCHOR.MIDDLE)
        if page:
            self.tbox(slide, SW - MARGIN - Inches(0.9), SRC_Y, Inches(0.9), SRC_H,
                      str(n), size=SZ_NOTE, color=fg, bold=True,
                      align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    def count_cover(self):
        """表紙など出典帯を置かないページでページ番号だけ進める"""
        self.page += 1

    def bullet_rows(self, slide, items, y=None, x=None, w=None, row_h=None,
                    gap=Inches(0.16), size=SZ_BODY, numbered=True,
                    fill=WHITE, line=BORDER, color=TEXT,
                    badge_fill=ORANGE, badge_color=NAVY, start=1):
        """縦積みの箇条書きカード（1行1メッセージ）"""
        x = MARGIN if x is None else x
        w = CW if w is None else w
        y = CY if y is None else y
        n = len(items)
        if row_h is None:
            row_h = int((CH - gap * (n - 1)) / n)
        for i, t in enumerate(items):
            self.rect(slide, x, y, w, row_h, fill=fill, line=line, radius=0.10)
            tx = x + Inches(0.35)
            if numbered:
                self.num_badge(slide, x + Inches(0.32),
                               y + (row_h - Inches(0.56)) / 2, start + i,
                               fill=badge_fill, color=badge_color)
                tx = x + Inches(1.12)
            self.tbox(slide, tx, y, w - (tx - x) - Inches(0.35), row_h, t,
                      size=size, color=color, bold=True,
                      anchor=MSO_ANCHOR.MIDDLE, ls=1.2)
            y += row_h + gap

    # ---------- 保存 ----------
    def save(self, pptx_path, refs_path=None):
        self.prs.save(pptx_path)
        if refs_path is None:
            refs_path = os.path.join(os.path.dirname(pptx_path), "slide_refs.json")
        with open(refs_path, "w", encoding="utf-8") as f:
            json.dump({"material": self.code, "version": self.version,
                       "slides": self.slide_refs}, f, ensure_ascii=False, indent=2)
        linked = sum(1 for v in self.slide_refs.values() if v["refs"])
        print("Saved:", pptx_path)
        print("Slides:", len(self.prs.slides))
        print("Saved:", refs_path,
              "(REF紐付け %d / 出典帯 %d ページ)" % (linked, len(self.slide_refs)))
