# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
def P(name):
    return os.path.join(HERE, name)

# ---- palette: PDPA official brand colors ----
# Navy #17375E / Orange #F5A623 / White #FFFFFF / Light Gray #F3F5F7
# (一般社団法人ペット防災アドバイザー協会 公式ブランドカラー)
NAVY_DARK   = RGBColor(0x17, 0x37, 0x5E)   # 公式メインカラー
NAVY_MED    = RGBColor(0x22, 0x45, 0x6A)   # ダーク背景上のカード用（メインより少し明るい）
NAVY_LINE   = RGBColor(0x35, 0x59, 0x7F)   # ダーク背景上の細い罫線用
ORANGE      = RGBColor(0xF5, 0xA6, 0x23)   # 公式サブカラー
ORANGE_DEEP = RGBColor(0xD9, 0x88, 0x12)   # 白背景上のテキスト用（サブより濃い）
CREAM       = RGBColor(0xF3, 0xF5, 0xF7)   # 公式補助カラー（ライトグレー）
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)   # 公式背景カラー
TEXT_DARK   = RGBColor(0x1F, 0x29, 0x37)
TEXT_MUTED  = RGBColor(0x76, 0x7C, 0x85)
CREAM_MUTED = RGBColor(0xC7, 0xCD, 0xD6)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def rect(slide, x, y, w, h, fill=None, line=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, x, y, w, h)
    if radius:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, text, size=18, color=TEXT_DARK, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.15,
            italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return tb


def add_pic_fill(slide, path, x, y, w, h):
    im = Image.open(path)
    iw, ih = im.size
    target_ratio = w / h
    src_ratio = iw / ih
    pic = slide.shapes.add_picture(path, x, y, width=w, height=h)
    if src_ratio > target_ratio:
        crop = (1 - target_ratio / src_ratio) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        crop = (1 - src_ratio / target_ratio) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    return pic


def rounded_pic(slide, path, x, y, w, h, radius=0.06):
    frame = rect(slide, x, y, w, h, fill=WHITE, radius=radius)
    add_pic_fill(slide, path, x, y, w, h)
    # rounded mask trick not natively supported for pictures in python-pptx; keep rectangular photo,
    # add a subtle border rectangle on top for a framed look
    border = rect(slide, x, y, w, h, fill=None, line=WHITE, radius=radius)
    border.line.width = Pt(2.5)
    return frame


def footer(slide, page, dark=False):
    color = CREAM_MUTED if dark else TEXT_MUTED
    textbox(slide, Inches(0.5), Inches(7.1), Inches(7), Inches(0.3),
            "一般社団法人 ペット防災アドバイザー協会　PDPA-101", size=9, color=color)
    textbox(slide, Inches(12.6), Inches(7.1), Inches(0.5), Inches(0.3),
            str(page), size=9, color=color, align=PP_ALIGN.RIGHT)


def kicker(slide, text, x=Inches(0.7), y=Inches(0.55), color=ORANGE):
    textbox(slide, x, y, Inches(5), Inches(0.4), text, size=14, color=color, bold=True)


def title(slide, text, x=Inches(0.7), y=Inches(0.9), size=30, color=NAVY_DARK, w=Inches(11.5)):
    textbox(slide, x, y, w, Inches(0.9), text, size=size, color=color, bold=True)
    rect(slide, x, y + Inches(0.85), Inches(0.9), Pt(3), fill=ORANGE)


def num_badge(slide, x, y, num, d=Inches(0.5), fill=ORANGE, color=WHITE):
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    circ.fill.solid(); circ.fill.fore_color.rgb = fill
    circ.line.fill.background(); circ.shadow.inherit = False
    tf = circ.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(num)
    r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = color; r.font.name = FONT
    return circ


# ============================================================
# Slide 1: 表紙
# ============================================================
s = add_slide()
set_bg(s, NAVY_DARK)
add_pic_fill(s, P("cover_senior_pet.jpg"), 0, 0, SW, SH)
overlay = rect(s, 0, Inches(3.6), SW, Inches(3.9), fill=NAVY_DARK)
overlay.fill.fore_color.rgb = NAVY_DARK
overlay.fill.transparency = 0
sp = overlay.fill._xPr.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
alpha = sp.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', {'val': '78000'})
sp.append(alpha)

logo = s.shapes.add_picture(P("logo.jpg"), Inches(0.7), Inches(0.55), height=Inches(1.1))
kicker(s, "第1回", x=Inches(0.7), y=Inches(4.0), color=ORANGE)
textbox(s, Inches(0.7), Inches(4.35), Inches(11), Inches(0.7),
        "ペット防災アドバイザー養成講座", size=34, color=WHITE, bold=True)
textbox(s, Inches(0.7), Inches(5.15), Inches(11), Inches(0.8),
        "ペット防災アドバイザーという新しい役割", size=24, color=ORANGE, bold=True)
textbox(s, Inches(0.7), Inches(6.6), Inches(9), Inches(0.4),
        "一般社団法人 ペット防災アドバイザー協会　公式教育プログラム　PDPA-101　Ver.2.0",
        size=11, color=CREAM_MUTED)

# ============================================================
# Slide 2: 学習目標
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "LEARNING GOALS")
title(s, "今日の学習目標")
items = [
    "ペット防災アドバイザーとは何か理解する",
    "日本で必要とされる理由を知る",
    "活動内容と、卒業後の活躍のイメージを持つ",
]
y = Inches(2.4)
for i, t in enumerate(items, start=1):
    rect(s, Inches(0.7), y, Inches(11.9), Inches(1.15), fill=WHITE, line=RGBColor(0xEA, 0xE6, 0xDE), radius=0.12)
    num_badge(s, Inches(1.05), y + Inches(0.32), i)
    textbox(s, Inches(1.9), y, Inches(10.4), Inches(1.15), t, size=19, color=TEXT_DARK, bold=True,
            anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.4)
footer(s, 2)

# ============================================================
# Slide 3: ペット防災アドバイザーとは
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "DEFINITION")
title(s, "ペット防災アドバイザーとは")
rounded_pic(s, P("advisor_role.jpg"), Inches(7.6), Inches(2.2), Inches(5.0), Inches(4.4))
rect(s, Inches(0.7), Inches(3.0), Inches(6.4), Inches(2.6), fill=NAVY_DARK, radius=0.06)
textbox(s, Inches(1.1), Inches(3.4), Inches(5.6), Inches(1.9),
        "「災害から人とペットの命を守るための\n備えを、地域に広める人」",
        size=24, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.5)
footer(s, 3)

# ============================================================
# Slide 4: 日本は世界有数の災害大国 (6 photos)
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "RISK IN JAPAN")
title(s, "日本は世界有数の災害大国")
disasters = [
    ("disaster_earthquake.jpg", "地震"),
    ("disaster_typhoon.jpg", "台風"),
    ("disaster_flood.jpg", "豪雨"),
    ("disaster_landslide.jpg", "土砂災害"),
    ("disaster_volcano.jpg", "火山"),
    ("disaster_snow.jpg", "豪雪"),
]
cell_w = Inches(1.9); gap = Inches(0.35)
total_w = cell_w * 6 + gap * 5
start_x = int((SW - total_w) / 2)
x = start_x
y = Inches(2.4)
for fname, label in disasters:
    rounded_pic(s, P(fname), x, y, cell_w, cell_w)
    textbox(s, x - Inches(0.15), y + cell_w + Inches(0.15), cell_w + Inches(0.3), Inches(0.4),
            label, size=15, color=NAVY_DARK, bold=True, align=PP_ALIGN.CENTER)
    x += cell_w + gap
footer(s, 4)

# ============================================================
# Slide 5: 飼い主が抱える不安
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "OWNER'S ANXIETY")
title(s, "飼い主が抱える不安")
rounded_pic(s, P("anxious_owner.jpg"), Inches(0.7), Inches(2.2), Inches(4.6), Inches(4.4))
items = [
    "避難所にペットを連れて行けるのか分からない",
    "何を備蓄すればいいのか分からない",
    "自分が動けなくなったらペットはどうなるのか不安",
]
y = Inches(2.3)
for i, t in enumerate(items, start=1):
    num_badge(s, Inches(5.75), y + Inches(0.1), i, fill=ORANGE_DEEP)
    textbox(s, Inches(6.5), y, Inches(6.1), Inches(1.1), t, size=18, color=TEXT_DARK, bold=True,
            anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.3)
    y += Inches(1.35)
footer(s, 5)

# ============================================================
# Slide 6: 同行避難 / 同伴避難
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "OFFICIAL GUIDELINE")
title(s, "環境省が示す「同行避難」の考え方")

card_w = Inches(5.75); card_h = Inches(4.5)
x1 = Inches(0.7); x2 = Inches(6.85); y0 = Inches(2.3)
rect(s, x1, y0, card_w, card_h, fill=WHITE, line=RGBColor(0xEA, 0xE6, 0xDE), radius=0.05)
rounded_pic(s, P("evacuation_walk.jpg"), x1, y0, card_w, Inches(2.1))
textbox(s, x1 + Inches(0.3), y0 + Inches(2.25), card_w - Inches(0.6), Inches(0.5), "同行避難",
        size=20, color=NAVY_DARK, bold=True)
textbox(s, x1 + Inches(0.3), y0 + Inches(2.75), card_w - Inches(0.6), Inches(1.3),
        "災害の発生時に、飼い主がペットを同行し、指定緊急避難場所等まで避難する『行動』そのものを指す。避難所等でペットと同室で飼養管理することを意味するものではない。",
        size=12.5, color=TEXT_DARK, line_spacing=1.3)
textbox(s, x1 + Inches(0.3), y0 + Inches(4.1), card_w - Inches(0.6), Inches(0.3),
        "出典：環境省『人とペットの災害対策ガイドライン』", size=9, color=TEXT_MUTED)

rect(s, x2, y0, card_w, card_h, fill=WHITE, line=RGBColor(0xEA, 0xE6, 0xDE), radius=0.05)
rounded_pic(s, P("shelter_pet_cage.jpg"), x2, y0, card_w, Inches(2.1))
textbox(s, x2 + Inches(0.3), y0 + Inches(2.25), card_w - Inches(0.6), Inches(0.5), "同伴避難",
        size=20, color=ORANGE_DEEP, bold=True)
textbox(s, x2 + Inches(0.3), y0 + Inches(2.75), card_w - Inches(0.6), Inches(1.3),
        "内閣府『避難所運営ガイドライン』で用いられる用語。避難所等において、被災者がペットを飼養管理する『状態』を指す。",
        size=12.5, color=TEXT_DARK, line_spacing=1.3)
textbox(s, x2 + Inches(0.3), y0 + Inches(4.1), card_w - Inches(0.6), Inches(0.3),
        "出典：内閣府『避難所運営ガイドライン』", size=9, color=TEXT_MUTED)
footer(s, 6)

# ============================================================
# Slide 7: アドバイザーの役割
# ============================================================
s = add_slide()
set_bg(s, NAVY_DARK)
kicker(s, "ROLE", color=ORANGE)
title(s, "アドバイザーの役割", color=WHITE)
roles = [
    ("正しい知識を伝え、生命を守る", "公的ガイドラインに基づく正確な情報を、わかりやすく伝える"),
    ("地域を支える防災リーダーになる", "自治体・地域活動と連携し、防災文化を広める"),
    ("飼い主に寄り添い、心配事を減らす", "断定的に指導するのではなく、一緒に備える伴走者として関わる"),
]
col_w = Inches(3.85); gap = Inches(0.25)
x = Inches(0.7)
for i, (h, d) in enumerate(roles, start=1):
    rect(s, x, Inches(2.4), col_w, Inches(4.0), fill=NAVY_MED, radius=0.08)
    num_badge(s, x + Inches(0.35), Inches(2.75), i)
    textbox(s, x + Inches(0.35), Inches(3.5), col_w - Inches(0.7), Inches(1.3), h,
            size=18, color=WHITE, bold=True, line_spacing=1.25)
    textbox(s, x + Inches(0.35), Inches(4.9), col_w - Inches(0.7), Inches(1.3), d,
            size=13, color=CREAM_MUTED, line_spacing=1.35)
    x += col_w + gap
footer(s, 7, dark=True)

# ============================================================
# Slide 8: 活動場所 (8 photos)
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "WHERE YOU CAN ACT")
title(s, "活動場所")
places = [
    ("location_townhall.jpg", "自治体"),
    ("location_school.jpg", "学校"),
    ("location_vetclinic.jpg", "動物病院"),
    ("location_grooming.jpg", "トリミングサロン"),
    ("location_pethotel.jpg", "ペットホテル"),
    ("location_shelter_org.jpg", "保護団体"),
    ("location_petshop.jpg", "ペットショップ"),
    ("location_community_event.jpg", "地域イベント"),
]
cell_w = Inches(2.75); cell_h = Inches(1.85); gap_x = Inches(0.2); gap_y = Inches(0.35)
cols = 4
total_w = cell_w * cols + gap_x * (cols - 1)
start_x = int((SW - total_w) / 2)
y0 = Inches(2.15)
for idx, (fname, label) in enumerate(places):
    col = idx % cols
    row = idx // cols
    x = start_x + col * (cell_w + gap_x)
    y = y0 + row * (cell_h + gap_y + Inches(0.3))
    rounded_pic(s, P(fname), x, y, cell_w, cell_h)
    ov = rect(s, x, y + cell_h - Inches(0.5), cell_w, Inches(0.5), fill=NAVY_DARK)
    textbox(s, x, y + cell_h - Inches(0.48), cell_w, Inches(0.45), label, size=13, color=WHITE,
            bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 8)

# ============================================================
# Slide 9: してはいけないこと
# ============================================================
s = add_slide()
set_bg(s, NAVY_DARK)
kicker(s, "BOUNDARIES", color=ORANGE)
title(s, "してはいけないこと", color=WHITE)
rect(s, Inches(0.7), Inches(2.3), Inches(11.9), Inches(4.2), fill=NAVY_MED, radius=0.05)
textbox(s, Inches(1.2), Inches(2.75), Inches(10.9), Inches(3.3),
        "アドバイザーの活動は「防災啓発・教育・情報提供」が中心です。\n\n"
        "獣医療行為、避難所運営の意思決定、行政の公式判断など、\n専門資格や権限が必要な領域には立ち入りません。\n\n"
        "判断に迷う相談を受けたときは、獣医師や自治体窓口など\n適切な専門職への「橋渡し役」であることを大切にしましょう。",
        size=18, color=WHITE, line_spacing=1.5)
footer(s, 9, dark=True)

# ============================================================
# Slide 10: 専門職との連携 (3 photos)
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "COLLABORATION")
title(s, "専門職との連携")
collabs = [
    ("collab_vet.jpg", "獣医師・看護師", "診療の合間に飼い主へ防災の声かけ"),
    ("collab_groomer.jpg", "トリマー・訓練士", "施術時に備蓄状況などをヒアリング"),
    ("collab_municipal.jpg", "自治体・保護団体", "地域の防災イベントや相談会で連携する"),
]
col_w = Inches(3.85); gap = Inches(0.25)
x = Inches(0.7)
for fname, h, d in collabs:
    rect(s, x, Inches(2.3), col_w, Inches(4.1), fill=WHITE, line=RGBColor(0xEA, 0xE6, 0xDE), radius=0.06)
    rounded_pic(s, P(fname), x, Inches(2.3), col_w, Inches(2.2))
    textbox(s, x + Inches(0.25), Inches(4.7), col_w - Inches(0.5), Inches(0.5), h,
            size=16, color=NAVY_DARK, bold=True)
    textbox(s, x + Inches(0.25), Inches(5.25), col_w - Inches(0.5), Inches(1.0), d,
            size=12, color=TEXT_MUTED, line_spacing=1.3)
    x += col_w + gap
footer(s, 10)

# ============================================================
# Slide 11: シニアアドバイザーの活動
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "SENIOR ADVISOR")
title(s, "シニアアドバイザーの活動")
items = [
    ("一般ペットオーナーへの啓発", "近隣のペットオーナーへの声かけや情報共有"),
    ("地域イベント・防災教室", "自治会・町内会での小規模な学習会の開催"),
    ("個別相談", "個々の家庭状況に応じた備えのアドバイス"),
]
y = Inches(2.4)
for i, (h, d) in enumerate(items, start=1):
    rect(s, Inches(0.7), y, Inches(11.9), Inches(1.25), fill=WHITE, line=RGBColor(0xEA, 0xE6, 0xDE), radius=0.1)
    num_badge(s, Inches(1.05), y + Inches(0.38), i)
    textbox(s, Inches(1.9), y + Inches(0.15), Inches(10.3), Inches(0.5), h, size=18, color=NAVY_DARK, bold=True)
    textbox(s, Inches(1.9), y + Inches(0.68), Inches(10.3), Inches(0.5), d, size=13, color=TEXT_MUTED)
    y += Inches(1.5)
footer(s, 11)

# ============================================================
# Slide 12: ペット関連従事者の活動
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "PET PROFESSIONALS")
title(s, "ペット関連従事者の活動")
items = [
    "自身の専門分野を核に活動する",
    "知り合いや関係業者（動物病院・ペットショップ等）と連携する",
    "より多くの飼い主へ、防災の重要性を伝えていく",
]
y = Inches(2.5)
for i, t in enumerate(items, start=1):
    num_badge(s, Inches(0.9), y, i, fill=ORANGE_DEEP)
    textbox(s, Inches(1.7), y - Inches(0.05), Inches(10.5), Inches(0.7), t, size=19, color=TEXT_DARK,
            bold=True, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.15)
footer(s, 12)

# ============================================================
# Slide 13: 認定制度の全体像
# ============================================================
s = add_slide()
set_bg(s, NAVY_DARK)
kicker(s, "CERTIFICATION", color=ORANGE)
title(s, "認定制度の全体像", color=WHITE)
textbox(s, Inches(0.7), Inches(2.3), Inches(4), Inches(0.5), "2段階", size=20, color=CREAM_MUTED)

box_w = Inches(4.2); box_h = Inches(1.3); y0 = Inches(3.2)
rect(s, Inches(1.0), y0, box_w, box_h, fill=NAVY_MED, line=ORANGE, radius=0.15)
textbox(s, Inches(1.0), y0, box_w, box_h, "認定ペット防災アドバイザー", size=18, color=WHITE, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.35), y0 + Inches(0.4), Inches(1.0), Inches(0.5))
arrow.fill.solid(); arrow.fill.fore_color.rgb = ORANGE; arrow.line.fill.background(); arrow.shadow.inherit = False
rect(s, Inches(6.5), y0, box_w, box_h, fill=ORANGE, radius=0.15)
textbox(s, Inches(6.5), y0, box_w, box_h, "認定講師", size=18, color=NAVY_DARK, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

textbox(s, Inches(1.0), Inches(5.1), Inches(9.7), Inches(0.8),
        "認定講師は、認定取得後1年以上の活動実績をもとに受験資格が付与される",
        size=14, color=CREAM_MUTED, line_spacing=1.4)
footer(s, 13, dark=True)

# ============================================================
# Slide 14: ケーススタディ①深夜の地震
# ============================================================
s = add_slide()
set_bg(s, NAVY_DARK)
add_pic_fill(s, P("case_study_night.jpg"), 0, 0, SW, SH)
ov = rect(s, 0, 0, SW, SH, fill=NAVY_DARK)
alpha_sp = ov.fill._xPr.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
alpha_el = alpha_sp.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}alpha', {'val': '62000'})
alpha_sp.append(alpha_el)

textbox(s, 0, Inches(0.9), SW, Inches(0.5), "ケーススタディ　①", size=18, color=ORANGE, bold=True,
        align=PP_ALIGN.CENTER)
textbox(s, 0, Inches(1.4), SW, Inches(0.9), "深夜の地震", size=36, color=WHITE, bold=True,
        align=PP_ALIGN.CENTER)
rect(s, Inches(2.15), Inches(2.7), Inches(9.0), Inches(2.6), fill=NAVY_DARK, radius=0.06)
textbox(s, Inches(2.55), Inches(3.0), Inches(8.2), Inches(2.0),
        "深夜2時、震度6強の地震が発生。\n犬はパニックで吠え続け、猫は押し入れに隠れています。\n\nあなたは最初に何をしますか？",
        size=20, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.5)
textbox(s, Inches(1.5), Inches(6.1), Inches(10.3), Inches(0.5),
        "あなたならどうしますか？（受講者に考えてもらう／正解を一つに決めず複数の視点を共有する）",
        size=11, color=CREAM_MUTED, align=PP_ALIGN.CENTER, italic=True)
footer(s, 14, dark=True)

# ============================================================
# Slide 15: ワーク
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "WORK - 5 MIN")
title(s, "ワーク：自分が活動したい地域を考える")
textbox(s, Inches(0.7), Inches(1.85), Inches(10), Inches(0.4),
        "以下を確認・記入してみましょう（5分）", size=13, color=TEXT_MUTED)
qs = [
    "あなたが将来アドバイザーとして活動したい地域・場所は？",
    "その地域には、どんな飼い主・ペットが多いと思いますか？",
    "まずはどんな小さな一歩から始められそうですか？",
]
y = Inches(2.55)
for i, q in enumerate(qs, start=1):
    rect(s, Inches(0.7), y, Inches(11.9), Inches(1.25), fill=WHITE, line=ORANGE, radius=0.1)
    textbox(s, Inches(1.1), y + Inches(0.15), Inches(0.8), Inches(0.5), f"Q{i}.", size=18, color=ORANGE_DEEP, bold=True)
    textbox(s, Inches(1.9), y, Inches(10.4), Inches(1.25), q, size=16, color=TEXT_DARK,
            anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.3)
    y += Inches(1.5)
footer(s, 15)

# ============================================================
# Slide 16: 今日のまとめ
# ============================================================
s = add_slide()
set_bg(s, NAVY_DARK)
kicker(s, "SUMMARY", color=ORANGE)
title(s, "今日のまとめ", color=WHITE)
textbox(s, Inches(0.7), Inches(1.95), Inches(6), Inches(0.4), "覚えること　3つ", size=15, color=CREAM_MUTED)
items = [
    "ペット防災アドバイザーは、防災啓発・教育・情報提供を通じて地域に貢献する存在",
    "「同行避難」と「同伴避難」は異なる意味を持つ、正確に理解する",
    "獣医療や行政判断が必要な相談は、専門職や自治体への橋渡しを行う",
]
y = Inches(2.5)
for i, t in enumerate(items, start=1):
    rect(s, Inches(0.7), y, Inches(11.9), Inches(1.25), fill=NAVY_MED, radius=0.1)
    num_badge(s, Inches(1.05), y + Inches(0.38), i)
    textbox(s, Inches(1.9), y, Inches(10.3), Inches(1.25), t, size=16, color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.3)
    y += Inches(1.5)
footer(s, 16, dark=True)

# ============================================================
# Slide 17: 理解度確認テスト
# ============================================================
s = add_slide()
set_bg(s, CREAM)
kicker(s, "CHECK YOUR UNDERSTANDING", x=Inches(4.5), y=Inches(1.6), color=ORANGE_DEEP)
tb = s.shapes.add_textbox(Inches(4.5), Inches(1.6), Inches(1), Inches(1))  # placeholder unused
textbox(s, 0, Inches(2.0), SW, Inches(0.9), "理解度確認テスト", size=34, color=NAVY_DARK, bold=True,
        align=PP_ALIGN.CENTER)
stats = [("10問", "問題数"), ("4択形式", "回答方式"), ("約10分", "所要時間")]
col_w = Inches(3.2); gap = Inches(0.4)
total_w = col_w * 3 + gap * 2
start_x = int((SW - total_w) / 2)
x = start_x
for val, label in stats:
    rect(s, x, Inches(3.3), col_w, Inches(1.8), fill=WHITE, line=RGBColor(0xEA, 0xE6, 0xDE), radius=0.1)
    textbox(s, x, Inches(3.55), col_w, Inches(0.9), val, size=30, color=ORANGE_DEEP, bold=True,
            align=PP_ALIGN.CENTER)
    textbox(s, x, Inches(4.5), col_w, Inches(0.4), label, size=13, color=TEXT_MUTED, align=PP_ALIGN.CENTER)
    x += col_w + gap
textbox(s, 0, Inches(5.6), SW, Inches(0.5), "今日の講義内容から出題します", size=15, color=TEXT_MUTED,
        align=PP_ALIGN.CENTER)
footer(s, 17)

# ============================================================
# Slide 18: 次回予告
# ============================================================
s = add_slide()
set_bg(s, NAVY_DARK)
kicker(s, "NEXT SESSION", color=ORANGE, x=Inches(0.7), y=Inches(2.6))
textbox(s, 0, Inches(3.0), SW, Inches(0.9), "次回予告", size=32, color=WHITE, bold=True,
        align=PP_ALIGN.CENTER)
textbox(s, 0, Inches(3.85), SW, Inches(0.7), "第2回　日本で起こる自然災害を知る", size=22, color=ORANGE,
        bold=True, align=PP_ALIGN.CENTER)
textbox(s, 0, Inches(4.8), SW, Inches(0.5), "お疲れさまでした。次回もよろしくお願いいたします。",
        size=15, color=CREAM_MUTED, align=PP_ALIGN.CENTER)
footer(s, 18, dark=True)

out_path = os.path.join(os.path.dirname(HERE), "PDPA-101_第1回スライド_redesigned.pptx")
prs.save(out_path)
print("Saved:", out_path)
