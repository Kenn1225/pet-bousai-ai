# -*- coding: utf-8 -*-
"""
認定ペット防災アドバイザー 活動収支計画書（Excel）

卒業後にセミナー開催・個別相談・講師派遣を行った場合の年間収支を試算する。
黄色いセルに数字を入れ替えれば、収入・支出・損益分岐点が自動で再計算される。

金額はすべて「仮置き」。実額に置き換えて使うことを前提にしている。
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "PDPA_activity_budget_plan.xlsx")

# ---- 協会ブランドカラー ----
NAVY = "17375E"
ORANGE = "F5A623"
LGRAY = "F3F5F7"
INPUT_BG = "FFF6DC"      # 入力セル（黄色）
BORDER_C = "D8DEE6"

FONT = "BIZ UDPGothic"

thin = Side(style="thin", color=BORDER_C)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()


def f(sz=10, bold=False, color="1B222C"):
    return Font(name=FONT, size=sz, bold=bold, color=color)


def put(ws, cell, value, size=10, bold=False, color="1B222C", bg=None,
        align=None, fmt=None, border=True, wrap=False):
    c = ws[cell]
    c.value = value
    c.font = f(size, bold, color)
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    else:
        c.alignment = Alignment(vertical="center")
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BOX
    return c


def section(ws, row, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    put(ws, "A%d" % row, text, size=11, bold=True, color="FFFFFF", bg=NAVY,
        align="left")
    ws.row_dimensions[row].height = 22


def header_row(ws, row, labels):
    for i, t in enumerate(labels, start=1):
        put(ws, "%s%d" % (get_column_letter(i), row), t, size=9, bold=True,
            color=NAVY, bg=LGRAY, align="center")


# ============================================================
# シート1：収支計画書
# ============================================================
ws = wb.active
ws.title = "収支計画書"
widths = [16, 30, 12, 10, 8, 14, 34]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.merge_cells("A1:G1")
put(ws, "A1", "認定ペット防災アドバイザー　活動収支計画書", size=16, bold=True,
    color="FFFFFF", bg=NAVY, align="center")
ws.row_dimensions[1].height = 32
ws.merge_cells("A2:G2")
put(ws, "A2", "一般社団法人 ペット防災アドバイザー協会　／　"
              "黄色のセルに数字を入れ替えると自動で再計算されます（金額はすべて仮置き）",
    size=9, color="5A6472", bg=LGRAY, align="center")

put(ws, "A3", "氏名", size=9, bold=True, bg=LGRAY, align="center")
put(ws, "B3", "", bg=INPUT_BG)
put(ws, "C3", "作成日", size=9, bold=True, bg=LGRAY, align="center")
ws.merge_cells("D3:E3")
put(ws, "D3", "", bg=INPUT_BG)
put(ws, "F3", "対象年度", size=9, bold=True, bg=LGRAY, align="center")
put(ws, "G3", "", bg=INPUT_BG)

# ---------------- 1. 前提条件 ----------------
r = 5
section(ws, r, "１．前提条件（ここだけ入力してください）")
r += 1
header_row(ws, r, ["区分", "項目", "単価（円）", "数量", "単位", "", "メモ"])
r += 1

PRE = [
    ("セミナー", "参加費（1人あたり）", 2000, None, "円", "地域相場に合わせて設定"),
    ("セミナー", "年間の開催回数", None, 12, "回", "月1回のペース"),
    ("セミナー", "1回あたりの参加者数", None, 10, "人", "公民館の小規模開催を想定"),
    ("個別相談", "相談料（1件あたり）", 5000, None, "円", "60分程度の想定"),
    ("個別相談", "年間の相談件数", None, 24, "件", "月2件のペース"),
    ("講師派遣", "謝金（1回あたり）", 30000, None, "円", "自治体・企業からの依頼"),
    ("講師派遣", "年間の依頼回数", None, 4, "回", "協会の紹介経由を含む"),
]
PRE_START = r
for kubun, item, tanka, suryo, tani, memo in PRE:
    put(ws, "A%d" % r, kubun, size=9, bold=True, color=NAVY, bg=LGRAY)
    put(ws, "B%d" % r, item, size=10)
    put(ws, "C%d" % r, tanka, bg=INPUT_BG if tanka is not None else None,
        fmt="#,##0", align="right")
    put(ws, "D%d" % r, suryo, bg=INPUT_BG if suryo is not None else None,
        fmt="#,##0", align="right")
    put(ws, "E%d" % r, tani, size=9, align="center")
    put(ws, "F%d" % r, "")
    put(ws, "G%d" % r, memo, size=9, color="5A6472")
    r += 1

P_FEE = "C%d" % PRE_START            # セミナー参加費
P_CNT = "D%d" % (PRE_START + 1)      # セミナー回数
P_PPL = "D%d" % (PRE_START + 2)      # 参加者数
S_FEE = "C%d" % (PRE_START + 3)      # 相談料
S_CNT = "D%d" % (PRE_START + 4)      # 相談件数
L_FEE = "C%d" % (PRE_START + 5)      # 講師謝金
L_CNT = "D%d" % (PRE_START + 6)      # 講師回数

# ---------------- 2. 収入の部 ----------------
r += 1
section(ws, r, "２．収入の部（年間）")
r += 1
header_row(ws, r, ["区分", "項目", "単価（円）", "数量", "単位", "金額（円）", "計算式"])
r += 1
inc_start = r
INC = [
    ("セミナー", "参加費収入", "=%s" % P_FEE, "=%s*%s" % (P_CNT, P_PPL), "人回",
     "参加費 × 開催回数 × 参加者数"),
    ("個別相談", "相談料収入", "=%s" % S_FEE, "=%s" % S_CNT, "件", "相談料 × 件数"),
    ("講師派遣", "謝金収入", "=%s" % L_FEE, "=%s" % L_CNT, "回", "謝金 × 依頼回数"),
]
for kubun, item, tanka, suryo, tani, keisan in INC:
    put(ws, "A%d" % r, kubun, size=9, bold=True, color=NAVY, bg=LGRAY)
    put(ws, "B%d" % r, item, size=10)
    put(ws, "C%d" % r, tanka, fmt="#,##0", align="right")
    put(ws, "D%d" % r, suryo, fmt="#,##0", align="right")
    put(ws, "E%d" % r, tani, size=9, align="center")
    put(ws, "F%d" % r, "=C%d*D%d" % (r, r), fmt="#,##0", align="right", bold=True)
    put(ws, "G%d" % r, keisan, size=9, color="5A6472")
    r += 1
inc_end = r - 1
INCOME_TOTAL = "F%d" % r
ws.merge_cells("A%d:E%d" % (r, r))
put(ws, "A%d" % r, "収入合計", size=11, bold=True, color="FFFFFF", bg=NAVY,
    align="right")
put(ws, "F%d" % r, "=SUM(F%d:F%d)" % (inc_start, inc_end), size=11, bold=True,
    color="FFFFFF", bg=NAVY, fmt="#,##0", align="right")
put(ws, "G%d" % r, "", bg=NAVY)
r += 2

# ---------------- 3. 初期費用 ----------------
section(ws, r, "３．支出の部　初期費用（初年度のみ）")
r += 1
header_row(ws, r, ["区分", "項目", "単価（円）", "数量", "単位", "金額（円）", "メモ"])
r += 1
init_start = r
INIT = [
    ("装備", "専用ベスト（ロゴ入り）", 6000, 1, "着", "協会オプション物品"),
    ("装備", "帽子（ロゴ入り）", 3000, 1, "個", "協会オプション物品"),
    ("装備", "腕章（ロゴ入り）", 2000, 1, "個", "協会オプション物品"),
    ("広報", "名刺作成", 3000, 1, "式", "100枚程度"),
    ("広報", "のぼり・バナー", 8000, 1, "式", "イベント出展用"),
    ("機材", "ポータブルスピーカー", 8000, 1, "台", "屋外・広い会場用"),
    ("機材", "展示用備品（クレート・備蓄品サンプル）", 15000, 1, "式", "実物を見せる用"),
]
for kubun, item, tanka, suryo, tani, memo in INIT:
    put(ws, "A%d" % r, kubun, size=9, bold=True, color=NAVY, bg=LGRAY)
    put(ws, "B%d" % r, item, size=10)
    put(ws, "C%d" % r, tanka, bg=INPUT_BG, fmt="#,##0", align="right")
    put(ws, "D%d" % r, suryo, bg=INPUT_BG, fmt="#,##0", align="right")
    put(ws, "E%d" % r, tani, size=9, align="center")
    put(ws, "F%d" % r, "=C%d*D%d" % (r, r), fmt="#,##0", align="right")
    put(ws, "G%d" % r, memo, size=9, color="5A6472")
    r += 1
init_end = r - 1
INIT_TOTAL = "F%d" % r
ws.merge_cells("A%d:E%d" % (r, r))
put(ws, "A%d" % r, "初期費用 小計", size=10, bold=True, color=NAVY, bg=LGRAY,
    align="right")
put(ws, "F%d" % r, "=SUM(F%d:F%d)" % (init_start, init_end), size=10, bold=True,
    color=NAVY, bg=LGRAY, fmt="#,##0", align="right")
put(ws, "G%d" % r, "", bg=LGRAY)
r += 2

# ---------------- 4. 固定費 ----------------
section(ws, r, "４．支出の部　固定費（毎年かかる）")
r += 1
header_row(ws, r, ["区分", "項目", "単価（円）", "数量", "単位", "金額（円）", "メモ"])
r += 1
fix_start = r
FIX = [
    ("保険", "団体賠償責任保険料", 15000, 1, "年", "★相見積が必要。仮置きの金額です"),
    ("広報", "チラシ・ポスター印刷", 12000, 1, "年", "年数回の刷り増し"),
    ("通信", "Zoom等の有料プラン", 2000, 12, "月", "オンライン開催をする場合"),
    ("事務", "郵送・事務用品", 6000, 1, "年", "資料送付など"),
]
for kubun, item, tanka, suryo, tani, memo in FIX:
    put(ws, "A%d" % r, kubun, size=9, bold=True, color=NAVY, bg=LGRAY)
    put(ws, "B%d" % r, item, size=10)
    put(ws, "C%d" % r, tanka, bg=INPUT_BG, fmt="#,##0", align="right")
    put(ws, "D%d" % r, suryo, bg=INPUT_BG, fmt="#,##0", align="right")
    put(ws, "E%d" % r, tani, size=9, align="center")
    put(ws, "F%d" % r, "=C%d*D%d" % (r, r), fmt="#,##0", align="right")
    put(ws, "G%d" % r, memo, size=9, color="5A6472")
    r += 1
fix_end = r - 1
FIX_TOTAL = "F%d" % r
ws.merge_cells("A%d:E%d" % (r, r))
put(ws, "A%d" % r, "固定費 小計", size=10, bold=True, color=NAVY, bg=LGRAY,
    align="right")
put(ws, "F%d" % r, "=SUM(F%d:F%d)" % (fix_start, fix_end), size=10, bold=True,
    color=NAVY, bg=LGRAY, fmt="#,##0", align="right")
put(ws, "G%d" % r, "", bg=LGRAY)
r += 2

# ---------------- 5. 変動費 ----------------
section(ws, r, "５．支出の部　変動費（開催・対応するたびにかかる）")
r += 1
header_row(ws, r, ["区分", "項目", "1回あたり（円）", "回数", "単位", "金額（円）", "メモ"])
r += 1

# セミナー1回あたりの内訳
sem_detail_start = r
SEM = [
    ("セミナー", "会場使用料（1回）", 2000, "公民館等。自治体で大きく異なります"),
    ("セミナー", "資料印刷（1人100円）", None, "参加者数に連動します"),
    ("セミナー", "交通費（1回）", 1000, "会場までの往復"),
    ("セミナー", "消耗品（1回）", 500, "文具・アンケート用紙など"),
]
for kubun, item, tanka, memo in SEM:
    put(ws, "A%d" % r, kubun, size=9, bold=True, color=NAVY, bg=LGRAY)
    put(ws, "B%d" % r, item, size=10)
    if tanka is None:
        put(ws, "C%d" % r, "=100*%s" % P_PPL, fmt="#,##0", align="right")
    else:
        put(ws, "C%d" % r, tanka, bg=INPUT_BG, fmt="#,##0", align="right")
    put(ws, "D%d" % r, "", size=9)
    put(ws, "E%d" % r, "円/回", size=9, align="center")
    put(ws, "F%d" % r, "", size=9)
    put(ws, "G%d" % r, memo, size=9, color="5A6472")
    r += 1
sem_detail_end = r - 1
SEM_UNIT = "C%d" % r
put(ws, "A%d" % r, "セミナー", size=9, bold=True, color=NAVY, bg=LGRAY)
put(ws, "B%d" % r, "1回あたり変動費 合計", size=10, bold=True)
put(ws, "C%d" % r, "=SUM(C%d:C%d)" % (sem_detail_start, sem_detail_end),
    fmt="#,##0", align="right", bold=True)
put(ws, "D%d" % r, "=%s" % P_CNT, fmt="#,##0", align="right")
put(ws, "E%d" % r, "回", size=9, align="center")
SEM_VAR = "F%d" % r
put(ws, "F%d" % r, "=C%d*D%d" % (r, r), fmt="#,##0", align="right", bold=True)
put(ws, "G%d" % r, "1回あたり × 開催回数", size=9, color="5A6472")
r += 1

put(ws, "A%d" % r, "個別相談", size=9, bold=True, color=NAVY, bg=LGRAY)
put(ws, "B%d" % r, "1件あたり変動費（交通費・資料）", size=10)
put(ws, "C%d" % r, 600, bg=INPUT_BG, fmt="#,##0", align="right")
put(ws, "D%d" % r, "=%s" % S_CNT, fmt="#,##0", align="right")
put(ws, "E%d" % r, "件", size=9, align="center")
put(ws, "F%d" % r, "=C%d*D%d" % (r, r), fmt="#,##0", align="right")
put(ws, "G%d" % r, "訪問しない場合は0でも可", size=9, color="5A6472")
var_row2 = r
r += 1

put(ws, "A%d" % r, "講師派遣", size=9, bold=True, color=NAVY, bg=LGRAY)
put(ws, "B%d" % r, "1回あたり変動費（交通費・資料）", size=10)
put(ws, "C%d" % r, 3500, bg=INPUT_BG, fmt="#,##0", align="right")
put(ws, "D%d" % r, "=%s" % L_CNT, fmt="#,##0", align="right")
put(ws, "E%d" % r, "回", size=9, align="center")
put(ws, "F%d" % r, "=C%d*D%d" % (r, r), fmt="#,##0", align="right")
put(ws, "G%d" % r, "遠方の場合は増額してください", size=9, color="5A6472")
var_row3 = r
r += 1

VAR_TOTAL = "F%d" % r
ws.merge_cells("A%d:E%d" % (r, r))
put(ws, "A%d" % r, "変動費 小計", size=10, bold=True, color=NAVY, bg=LGRAY,
    align="right")
put(ws, "F%d" % r, "=%s+F%d+F%d" % (SEM_VAR, var_row2, var_row3), size=10,
    bold=True, color=NAVY, bg=LGRAY, fmt="#,##0", align="right")
put(ws, "G%d" % r, "", bg=LGRAY)
r += 2

# ---------------- 6. 収支まとめ ----------------
section(ws, r, "６．収支のまとめ")
r += 1
header_row(ws, r, ["", "項目", "", "", "", "金額（円）", "内容"])
r += 1
SUM_ROWS = [
    ("収入合計", "=%s" % INCOME_TOTAL, "セミナー＋個別相談＋講師派遣", False),
    ("支出合計（初年度）", "=%s+%s+%s" % (INIT_TOTAL, FIX_TOTAL, VAR_TOTAL),
     "初期費用＋固定費＋変動費", False),
    ("初年度の収支", None, "装備をそろえる年", True),
    ("支出合計（2年目以降）", "=%s+%s" % (FIX_TOTAL, VAR_TOTAL),
     "固定費＋変動費", False),
    ("2年目以降の収支", None, "初期費用がなくなります", True),
]
first_sum = r
for i, (label, formula, memo, emph) in enumerate(SUM_ROWS):
    ws.merge_cells("B%d:E%d" % (r, r))
    put(ws, "A%d" % r, "", bg=LGRAY if emph else None)
    put(ws, "B%d" % r, label, size=11 if emph else 10, bold=True,
        color="FFFFFF" if emph else "1B222C", bg=ORANGE if emph else None)
    for cc in "CDE":
        ws["%s%d" % (cc, r)].fill = PatternFill("solid", fgColor=ORANGE) if emph else PatternFill()
        ws["%s%d" % (cc, r)].border = BOX
    if formula is None:
        if "初年度" in label:
            formula = "=F%d-F%d" % (first_sum, first_sum + 1)
        else:
            formula = "=F%d-F%d" % (first_sum, first_sum + 3)
    put(ws, "F%d" % r, formula, size=12 if emph else 10, bold=True,
        color="FFFFFF" if emph else "1B222C", bg=ORANGE if emph else None,
        fmt="#,##0", align="right")
    put(ws, "G%d" % r, memo, size=9, color="5A6472")
    r += 1
r += 1

# ---------------- 7. 損益分岐点 ----------------
section(ws, r, "７．損益分岐点（何回開催すれば赤字にならないか）")
r += 1
header_row(ws, r, ["", "項目", "", "", "", "数値", "計算の考え方"])
r += 1
put(ws, "B%d" % r, "セミナー1回あたりの粗利", size=10)
ws.merge_cells("B%d:E%d" % (r, r))
BEP_UNIT = "F%d" % r
put(ws, "F%d" % r, "=%s*%s-%s" % (P_FEE, P_PPL, SEM_UNIT), fmt="#,##0",
    align="right", bold=True)
put(ws, "G%d" % r, "参加費×参加者数 − 1回あたり変動費", size=9, color="5A6472")
r += 1
ws.merge_cells("B%d:E%d" % (r, r))
put(ws, "B%d" % r, "固定費を回収するのに必要な開催回数", size=10)
put(ws, "F%d" % r, '=IF(%s<=0,"—",ROUNDUP(%s/%s,0))' % (BEP_UNIT, FIX_TOTAL, BEP_UNIT),
    fmt="#,##0", align="right", bold=True)
put(ws, "G%d" % r, "固定費 ÷ 1回あたりの粗利", size=9, color="5A6472")
r += 1
ws.merge_cells("B%d:E%d" % (r, r))
put(ws, "B%d" % r, "初年度に初期費用まで回収する開催回数", size=10)
put(ws, "F%d" % r, '=IF(%s<=0,"—",ROUNDUP((%s+%s)/%s,0))'
    % (BEP_UNIT, FIX_TOTAL, INIT_TOTAL, BEP_UNIT), fmt="#,##0", align="right",
    bold=True)
put(ws, "G%d" % r, "（固定費＋初期費用）÷ 1回あたりの粗利", size=9, color="5A6472")
r += 2

ws.merge_cells("A%d:G%d" % (r, r))
put(ws, "A%d" % r, "※ 金額はすべて仮置きです。実際の会場費・保険料・参加費に置き換えてお使いください。"
                   "詳しい注意点は「前提と注意」シートをご覧ください。",
    size=9, color="C02A1E", bg=LGRAY)

ws.freeze_panes = "A4"
ws.sheet_view.showGridLines = False

# ============================================================
# シート2：シナリオ比較
# ============================================================
ws2 = wb.create_sheet("シナリオ比較")
for i, w in enumerate([26, 16, 16, 16, 34], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.sheet_view.showGridLines = False

ws2.merge_cells("A1:E1")
put(ws2, "A1", "活動ペース別の年間収支シミュレーション", size=14, bold=True,
    color="FFFFFF", bg=NAVY, align="center")
ws2.row_dimensions[1].height = 28
ws2.merge_cells("A2:E2")
put(ws2, "A2", "「標準」は収支計画書シートと同じ前提です。黄色のセルは入力できます。",
    size=9, color="5A6472", bg=LGRAY, align="center")

r2 = 4
header_row(ws2, r2, ["項目", "控えめ", "標準", "積極的", "メモ"])
r2 += 1
SCEN = [
    ("セミナー 年間開催回数", 6, 12, 24, "回"),
    ("セミナー 1回の参加者数", 8, 10, 12, "人"),
    ("セミナー 参加費", 1500, 2000, 2500, "円/人"),
    ("個別相談 年間件数", 12, 24, 48, "件"),
    ("個別相談 相談料", 3000, 5000, 5000, "円/件"),
    ("講師派遣 年間回数", 0, 4, 8, "回"),
    ("講師派遣 謝金", 30000, 30000, 30000, "円/回"),
]
sc_start = r2
for name, a, b, c, memo in SCEN:
    put(ws2, "A%d" % r2, name, size=10)
    for col, v in zip("BCD", (a, b, c)):
        put(ws2, "%s%d" % (col, r2), v, bg=INPUT_BG, fmt="#,##0", align="right")
    put(ws2, "E%d" % r2, memo, size=9, color="5A6472")
    r2 += 1
sc = {n: sc_start + i for i, (n, *_ ) in enumerate(SCEN)}
r2 += 1

header_row(ws2, r2, ["", "控えめ", "標準", "積極的", ""])
r2 += 1
put(ws2, "A%d" % r2, "年間収入", size=10, bold=True)
for col in "BCD":
    put(ws2, "%s%d" % (col, r2),
        "={c}{p}*{c}{n}*{c}{f}+{c}{sc}*{c}{sf}+{c}{lc}*{c}{lf}".format(
            c=col, p=sc["セミナー 年間開催回数"], n=sc["セミナー 1回の参加者数"],
            f=sc["セミナー 参加費"], sc=sc["個別相談 年間件数"],
            sf=sc["個別相談 相談料"], lc=sc["講師派遣 年間回数"],
            lf=sc["講師派遣 謝金"]),
        fmt="#,##0", align="right", bold=True)
put(ws2, "E%d" % r2, "参加費＋相談料＋謝金", size=9, color="5A6472")
inc_row = r2
r2 += 1

put(ws2, "A%d" % r2, "年間支出（2年目以降）", size=10, bold=True)
for col in "BCD":
    put(ws2, "%s%d" % (col, r2),
        "=収支計画書!{fix}+({c}{p}*4500)+({c}{sc}*600)+({c}{lc}*3500)".format(
            fix=FIX_TOTAL, c=col, p=sc["セミナー 年間開催回数"],
            sc=sc["個別相談 年間件数"], lc=sc["講師派遣 年間回数"]),
        fmt="#,##0", align="right")
put(ws2, "E%d" % r2, "固定費＋変動費（1回4,500円で概算）", size=9, color="5A6472")
exp_row = r2
r2 += 1

put(ws2, "A%d" % r2, "年間収支", size=11, bold=True, color="FFFFFF", bg=ORANGE)
for col in "BCD":
    put(ws2, "%s%d" % (col, r2), "={c}{i}-{c}{e}".format(c=col, i=inc_row, e=exp_row),
        size=12, bold=True, color="FFFFFF", bg=ORANGE, fmt="#,##0", align="right")
put(ws2, "E%d" % r2, "", bg=ORANGE)
r2 += 2

ws2.merge_cells("A%d:E%d" % (r2, r2))
put(ws2, "A%d" % r2,
    "※ 変動費は概算です。正確に見るときは「収支計画書」シートで1回あたりの内訳を入れてください。",
    size=9, color="C02A1E", bg=LGRAY)

# ============================================================
# シート3：前提と注意
# ============================================================
ws3 = wb.create_sheet("前提と注意")
for i, w in enumerate([4, 30, 78], start=1):
    ws3.column_dimensions[get_column_letter(i)].width = w
ws3.sheet_view.showGridLines = False

ws3.merge_cells("A1:C1")
put(ws3, "A1", "この計画書を使う前に確認すること", size=14, bold=True,
    color="FFFFFF", bg=NAVY, align="center")
ws3.row_dimensions[1].height = 28

NOTES = [
    ("金額の扱い",
     "本シートの金額はすべて仮置きです。ご自身の地域・活動内容の実額に置き換えてください。"),
    ("会場使用料",
     "自治体の公民館・地区センターは無料〜数千円と幅があります。減免制度がある自治体も多いので、"
     "事前に窓口へご確認ください。民間施設を借りる場合は大きく変わります。"),
    ("賠償責任保険料",
     "★未確定です。損害保険会社への相見積が必要です。"
     "協会の条件シート（pdpa_insurance_rfq.html）をご利用ください。"
     "保険料は認定者負担で、協会は手数料を上乗せしません。"),
    ("参加費の設定",
     "地域の相場と、自治体主催（無料開催が多い）か自主開催かで大きく変わります。"
     "自治体からの依頼は参加費ではなく謝金という形になるのが一般的です。"),
    ("消費税",
     "課税売上高が1,000万円以下であれば免税事業者となるのが原則です。"
     "取引先からインボイス（適格請求書）を求められる場合は登録の検討が必要です。"),
    ("所得税・確定申告",
     "給与を1か所から受けていて源泉徴収されている方は、給与所得・退職所得以外の所得の合計額が"
     "20万円を超えると確定申告が必要です（国税庁 タックスアンサー No.1900）。"
     "開業届や青色申告の要否を含め、最寄りの税務署または税理士にご確認ください。"),
    ("活動範囲",
     "認定者の活動は防災啓発・教育・情報提供が中心です。獣医療行為、避難所運営の意思決定、"
     "行政の公式判断は行いません。判断に迷う相談は獣医師や自治体窓口へお繋ぎください。"),
    ("協会からの支援",
     "セミナー開催キット、協会名義の公式文書、認定者マップからの紹介、共同開催マッチングを"
     "利用すると、広報費や準備の負担を下げられます。"),
]
r3 = 3
for i, (title, body) in enumerate(NOTES, start=1):
    put(ws3, "A%d" % r3, i, size=10, bold=True, color="FFFFFF", bg=ORANGE,
        align="center")
    put(ws3, "B%d" % r3, title, size=11, bold=True, color=NAVY, bg=LGRAY)
    put(ws3, "C%d" % r3, body, size=10, wrap=True, align="left")
    ws3.row_dimensions[r3].height = 46
    r3 += 1

r3 += 1
ws3.merge_cells("A%d:C%d" % (r3, r3))
put(ws3, "A%d" % r3, "© 一般社団法人 ペット防災アドバイザー協会　"
                     "教材コード PDPA-100S 関連資料　Ver.1.0",
    size=9, color="5A6472", bg=LGRAY, align="center")

wb.save(OUT)
print("Saved:", OUT)
